"""v36 -> v37: paper-style figure on slide 12; slide 13 slimmed, speech enriched.

Slide 12: the figure is replaced by assets/slide12_graph_hypergraph.png as
redrawn by fig_slide12.py in the visual grammar of the recommender-systems
literature: panel (a) the bipartite graph G = (U, I, E) with the 2-hop message
as a bold directed path and the LightGCN / HCCF / HPCF symmetric-normalisation
rule set in mathematical type; panel (b) the hypergraph H = (V, E) with dashed
hyperedge contours, a context node and the incidence matrix H beside them.

Slide 13: the merged problem-statement slide carried full sentences in seven
cards plus the gap bar.  The slide now keeps only titles and one-line keys;
everything that was removed moves into the spoken note, which grows from 103
to 190 words and carries the full argument (sparsity, cold-start, the
filter-bubble loop, interpretability, local vs global clustering explanations,
the three structural limits with their details, the gap and the claim).
"""
import hashlib
import json
import os

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v36.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v37.pptx")
FIG = os.path.join(HERE, "assets", "slide12_graph_hypergraph.png")

CREAM = "F1EDE6"
INK = "111111"

ROW1_BODY = {
    29: "Almost-empty interaction matrix.",
    30: "No history for new users or items.",
    31: "Exposure feeds exposure.",
    32: "The limit this thesis targets.",
}
ROW2_BODY = {
    36: [("Explanations are neither ", False), ("faithful nor actionable", True), (".", False)],
    40: [("Local explanations break ", False), ("across levels and at scale", True), (".", False)],
    44: [("Post-hoc", True), (" only: training never sees them.", False)],
}
GAP = [("No ", False), ("cooperative-attribution framework", True),
       (" explains clustering faithfully, stays consistent across levels, and works as an ", False),
       ("in-training signal", True), (". ", False), ("Claim:", True),
       (" Shapley-value attribution can be that framework.", False)]

NOTE_13 = (
    "Four classical limits appear from this evolution. Data sparsity: the user-item matrix is almost "
    "empty, so there is very little signal to learn from. Cold-start: new users and items arrive with "
    "no history. Popularity bias: exposure leads to interaction, which leads to more exposure, a "
    "filter-bubble loop that also squeezes diversity. And above all, the absence of interpretability: "
    "the limit this thesis targets. For clustering specifically, methods give a local or a global "
    "explanation, not both, and rarely stay consistent across levels of detail. Three structural "
    "limitations follow. One, complex models are still hard to explain in a way that is faithful and "
    "actionable. Two, local explanations do not carry over to multi-level structures or large datasets: "
    "a method that works on a toy partition may break on hundreds of thousands of nested records. Three, "
    "most explanations stay post-hoc: they never shape how the model learns, nor how it handles the "
    "accuracy, diversity and context trade-off. The gap bar closes: the literature still lacks one "
    "cooperative-attribution framework that explains clustering faithfully, stays consistent across "
    "levels, and then works as an in-training signal in recommendation. Claim: Shapley-value attribution "
    "can be that framework."
)


def main():
    prs = Presentation(SRC)

    # ---- slide 12: swap the picture part
    s12 = prs.slides[11]
    pic = [sh for sh in s12.shapes if sh.shape_id == 33][0]
    rId = pic._element.blipFill.blip.get(qn("r:embed"))
    part = s12.part.related_part(rId)
    part._blob = open(FIG, "rb").read()

    # ---- slide 13: slim the cards, enrich the note
    s13 = prs.slides[12]
    by_id = {sh.shape_id: sh for sh in s13.shapes}
    for sid, text in ROW1_BODY.items():
        tf = by_id[sid].text_frame
        para = tf.paragraphs[2]
        for r in list(para.runs)[1:]:
            r._r.getparent().remove(r._r)
        para.runs[0].text = text
        para.runs[0].font.size = Pt(20)
        tf.paragraphs[0].runs[0].font.size = Pt(22)
        tf.paragraphs[1].runs[0].font.size = Pt(20)
    for sid, runs in ROW2_BODY.items():
        tf = by_id[sid].text_frame
        for p in list(tf.paragraphs):
            p._p.getparent().remove(p._p)
        p = tf.add_paragraph()
        for text, bold in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(20)
            r.font.bold = bold
            r.font.name = "Nunito Semi-Bold"
            from pptx.dml.color import RGBColor
            r.font.color.rgb = RGBColor(0xF1, 0xED, 0xE6)
    for sid in (35, 39, 43):
        by_id[sid].text_frame.paragraphs[0].runs[0].font.size = Pt(22)
    gtf = by_id[47].text_frame
    for p in list(gtf.paragraphs):
        p._p.getparent().remove(p._p)
    p = gtf.add_paragraph()
    from pptx.dml.color import RGBColor
    for text, bold in GAP:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(18)
        r.font.bold = bold
        r.font.name = "Nunito Semi-Bold"
        r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    s13.notes_slide.notes_text_frame.text = NOTE_13

    prs.save(DST)

    back = Presentation(DST)
    nj = {str(i): (sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else "")
          for i, sl in enumerate(back.slides, 1)}
    json.dump(nj, open(os.path.join(HERE, "notes_v37.json"), "w"), indent=1, ensure_ascii=False)
    old = Presentation(SRC)
    dnotes = [i for i, (a, b) in enumerate(zip(
        [s.notes_slide.notes_text_frame.text for s in old.slides],
        [s.notes_slide.notes_text_frame.text for s in back.slides]), 1) if a != b]
    dimg = [i for i, (a, b) in enumerate(zip(old.slides, back.slides), 1)
            if [sh.image.sha1 for sh in a.shapes if sh.shape_type == 13] !=
               [sh.image.sha1 for sh in b.shapes if sh.shape_type == 13]]
    words = sum(len(nj[str(i)].split()) for i in range(1, 76))
    print("notes changed:", dnotes, "| images changed:", dimg)
    print("slides:", len(back.slides))
    print("spoken words 1-75:", words, f"= {words/130:.1f} min @130, {words/120:.1f} @120")
    print("note 13 words:", len(nj["13"].split()))
    a = hashlib.sha1(open(FIG, "rb").read()).hexdigest()
    pic2 = [sh for sh in back.slides[11].shapes if sh.shape_id == 33][0]
    print("fig embedded:", pic2.image.sha1 == a)


if __name__ == "__main__":
    main()
