#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# patch_v32.py: v31 -> v32   -- slide 4 describes, it no longer asks
#
# Slide 4 is the first slide where the talk starts describing the subject, so it
# should state things rather than put three questions on the screen. Same three
# cards, same layout, same colours; the content changes from a question + hint to
# a claim + the facts behind it, and the headline says what the slide is about.
#
# Speech matches the screen one to one and stays in spoken register.
# ---------------------------------------------------------------------------
import json
import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v31.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v32.pptx")
NOTES_IN = os.path.join(HERE, "notes_v31.json")
NOTES_OUT = os.path.join(HERE, "notes_v32.json")

SLIDE = 4
OLD_TITLE = "Motivation: Three Questions"
NEW_TITLE = "Motivation: Why Explainability Matters"

CARDS = {
    "TextBox 30": (
        "Recommenders decide at scale.",
        "They shape what billions of users see, buy and watch every day: news, study, health and credit decisions.",
    ),
    "TextBox 33": (
        "The reasoning behind those decisions is hidden.",
        "Matrix factorisation hid it in latent factors; deep and graph models hide it in message passing. Neither can be audited or acted on.",
    ),
    "TextBox 36": (
        "Transparency is now expected.",
        "Users, designers and regulators now ask for reasons. It has to be built into the model, not added afterwards.",
    ),
}

NOTE_4 = (
    "Let me start with why this matters.\n\n"
    "Recommenders are everywhere: they already decide much of what billions of "
    "people see, buy and watch, from news to health and credit. And the "
    "reasoning behind those decisions is hidden, in latent factors or in "
    "message passing, so nothing can be audited or acted on.\n\n"
    "Users, designers and regulators all ask for reasons now. So transparency "
    "has to be built into the model, not added afterwards. This thesis treats "
    "accuracy and interpretability as goals to be met together."
)


def set_paragraph(paragraph, text):
    """Rewrite a paragraph, keeping the formatting of its first run."""
    runs = paragraph.runs
    if not runs:
        return False
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    return True


def main():
    prs = Presentation(SRC)
    slide = prs.slides[SLIDE - 1]
    report = []

    # --- headline ----------------------------------------------------------
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == OLD_TITLE:
            set_paragraph(sh.text_frame.paragraphs[0], NEW_TITLE)
            report.append(f"headline: {OLD_TITLE!r} -> {NEW_TITLE!r}")
            break
    else:
        report.append(f"!! headline {OLD_TITLE!r} not found")

    # --- the three cards ---------------------------------------------------
    for name, (claim, facts) in CARDS.items():
        shape = next((sh for sh in slide.shapes if sh.name == name), None)
        if shape is None or not shape.has_text_frame:
            report.append(f"!! {name} not found")
            continue
        paras = shape.text_frame.paragraphs
        set_paragraph(paras[0], claim)
        set_paragraph(paras[1], facts)
        report.append(f"{name}: {len(claim.split())}+{len(facts.split())} words")

    # --- the speech --------------------------------------------------------
    notes = json.load(open(NOTES_IN))
    before = len(notes[str(SLIDE)].split())
    notes[str(SLIDE)] = NOTE_4
    slide.notes_slide.notes_text_frame.text = NOTE_4
    json.dump(notes, open(NOTES_OUT, "w"), ensure_ascii=False, indent=1)
    report.append(f"note {SLIDE}: {before} -> {len(NOTE_4.split())} words")

    prs.save(DST)
    spoken = sum(len(notes[str(i)].split()) for i in range(1, 77))
    report.append(f"wrote {os.path.relpath(DST, HERE)}")
    report.append(f"wrote {os.path.relpath(NOTES_OUT, HERE)}")
    report.append(f"spoken words slides 1-76: {spoken} "
                  f"({spoken / 130:.1f} min @130 wpm, {spoken / 120:.1f} @120 wpm)")
    print("\n".join(report))


if __name__ == "__main__":
    main()
