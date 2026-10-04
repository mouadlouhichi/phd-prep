#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# patch_v30.py: v29 -> v30   -- slide 6 keeps two things only
#
# Slide 6 had four blocks: the evolution timeline, a one-line conclusion about
# the growing gap, the card "Why the gap matters" and the card "What this thesis
# argues". The candidate asked for the timeline and "Why the gap matters" only,
# so the other two blocks go:
#
#   removed   TextBox 39  "The interpretability gap grows: ..."
#             Rounded Rectangle 43 / 44 / TextBox 45   "What this thesis argues"
#   kept      the timeline (label, five stages, four arrows, the rule)
#             the card "Why the gap matters", now full width, top at the same
#             distance under the rule as before, bottom clear of the citation band
#   laid out  the card's three bullets become three columns inside the card, which
#             is what makes a full-width card look deliberate; each column is a
#             copy of the existing bullet paragraph, so the bold lead-in, the
#             yellow bullet and the white text keep their exact formatting
#
# The speech for slide 6 loses its last paragraph ("So what does this thesis
# argue? ..."), which was the part about the removed card.
# ---------------------------------------------------------------------------
import json
import os
from copy import deepcopy

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v29.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v30.pptx")
NOTES_IN = os.path.join(HERE, "notes_v28.json")
NOTES_OUT = os.path.join(HERE, "notes_v30.json")

EMU_IN = 914400
SLIDE = 6

# the band on this slide is three lines: its text starts at y = 9.80
CONTENT_BOTTOM = 9.74
CARD_TOP = 6.38            # 0.6 in under the rule, as the deck had it before
CARD_LEFT, CARD_WIDTH = 1.12, 17.76
COL_TOP, COL_HEIGHT = 6.72, 2.90
BULLET_SIZE = 22.0         # the deck's card bullets are 19 pt; at 22 pt the three lines fill this taller card
COL_WIDTH, COL_GAP = 5.32, 0.4

DROP = ("TextBox 39", "Rounded Rectangle 43", "Rounded Rectangle 44", "TextBox 45")

NOTE_6 = (
    "Let me put this in context.\n\n"
    "Recommendation systems started with similarity models, then matrix "
    "factorisation, then neural and graph models, and now hypergraphs. Each "
    "step made the ranking stronger, and each step hid more of the reasoning. "
    "So the gap keeps growing: the models do more, and we can explain less.\n\n"
    "Why does this matter? Three reasons. Users get outputs without any "
    "reason, so trust drops. Designers cannot debug what they cannot inspect. "
    "And regulators now ask for transparency."
)


def sp_of(shape):
    return shape._element.find(qn("p:spPr"))


def set_box(shape, left, top, width, height):
    shape.left, shape.top = Emu(int(left * EMU_IN)), Emu(int(top * EMU_IN))
    shape.width, shape.height = Emu(int(width * EMU_IN)), Emu(int(height * EMU_IN))


def centre_text(shape):
    bodyPr = shape._element.find(qn("p:txBody")).find(qn("a:bodyPr"))
    bodyPr.set("anchor", "ctr")


def only_paragraph(sp_element, index):
    """Keep just one <a:p> of a copied shape, so a column shows one bullet."""
    txBody = sp_element.find(qn("p:txBody"))
    paras = txBody.findall(qn("a:p"))
    keep = paras[index]
    for p in paras:
        if p is not keep:
            txBody.remove(p)


def main():
    prs = Presentation(SRC)
    slide = prs.slides[SLIDE - 1]
    shapes = {sh.name: sh for sh in slide.shapes}
    report = []

    # --- what goes ---------------------------------------------------------
    for name in DROP:
        if name not in shapes:
            report.append(f"!! {name} not found")
            continue
        shapes[name]._element.getparent().remove(shapes[name]._element)
        report.append(f"removed {name}")

    # --- the kept card, full width ----------------------------------------
    panel, chip, bullets = shapes["Rounded Rectangle 40"], shapes["Rounded Rectangle 41"], shapes["TextBox 42"]
    set_box(panel, CARD_LEFT, CARD_TOP, CARD_WIDTH, CONTENT_BOTTOM - CARD_TOP)
    set_box(chip, 1.62, 6.12, Emu(chip.width).inches, Emu(chip.height).inches)
    set_box(bullets, 1.62, COL_TOP, COL_WIDTH, COL_HEIGHT)
    report.append(f"card {CARD_LEFT}-{CARD_LEFT + CARD_WIDTH} in, y {CARD_TOP}-{CONTENT_BOTTOM}")

    # --- the three bullets as three columns --------------------------------
    txBody = bullets._element.find(qn("p:txBody"))
    if len(txBody.findall(qn("a:p"))) != 3:
        report.append("!! the card no longer holds three bullets")
    prev = bullets._element
    for i in range(1, 3):
        copy = deepcopy(bullets._element)
        only_paragraph(copy, i)
        prev.addnext(copy)                      # keep the columns in reading order
        prev = copy
        new_shape = [sh for sh in slide.shapes if sh._element is copy][0]
        set_box(new_shape, 1.62 + i * (COL_WIDTH + COL_GAP), COL_TOP, COL_WIDTH, COL_HEIGHT)
        centre_text(new_shape)
        report.append(f"column {i + 1} at x {1.62 + i * (COL_WIDTH + COL_GAP):.2f}, one bullet")
    only_paragraph(bullets._element, 0)
    for sh in slide.shapes:                      # a card bullet is 19 pt in this deck
        if sh.name != "TextBox 42":
            continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(BULLET_SIZE)
    centre_text(bullets)

    # --- the speech --------------------------------------------------------
    notes = json.load(open(NOTES_IN))
    notes[str(SLIDE)] = NOTE_6
    slide.notes_slide.notes_text_frame.text = NOTE_6      # into the deck and the json
    json.dump(notes, open(NOTES_OUT, "w"), ensure_ascii=False, indent=1)
    report.append(f"note {SLIDE}: rewritten without the removed card "
                  f"({len(NOTE_6.split())} words)")
    spoken = sum(len(notes[str(i)].split()) for i in range(1, 77))
    report.append(f"wrote {os.path.relpath(DST, HERE)}")
    report.append(f"wrote {os.path.relpath(NOTES_OUT, HERE)}")
    report.append(f"spoken words slides 1-76: {spoken} "
                  f"({spoken / 130:.1f} min @130 wpm, {spoken / 120:.1f} @120 wpm)")

    prs.save(DST)
    print("\n".join(report))


if __name__ == "__main__":
    main()
