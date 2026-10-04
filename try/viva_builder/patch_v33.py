#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# patch_v33.py: v32 -> v33
#
# One change: the speech on slide 16, built from the candidate's own text:
#
#   "These research questions lead to three main concrete contributions, which
#    I'll present in order ... The third contribution, A Dynamic Hypergraph
#    Cooperative Game for Preference-aware Recommendation ... Together, these
#    three contributions follow a common progression: from explanation, to
#    scalability, and finally to action."
#
# The third sentence was a fragment (it named the work but never said what it
# does), and the first was missing its full stop. The finished note keeps the
# candidate's wording and progression, completes the third contribution, and
# names each contribution's ground, which is what the three cards on the slide
# show (Wine Quality, Beijing Air Quality, MovieLens-1M and Amazon-Book).
# ---------------------------------------------------------------------------
import json
import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v32.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v33.pptx")
NOTES_IN = os.path.join(HERE, "notes_v32.json")
NOTES_OUT = os.path.join(HERE, "notes_v33.json")

SLIDE = 16

NOTE_16 = (
    "These research questions lead to three concrete contributions, which I "
    "will present in order.\n\n"
    "The first contribution introduces a Shapley-based framework for "
    "explaining black-box clustering, demonstrated on the Wine Quality data. "
    "The second contribution extends that idea to large-scale, multi-level "
    "clustering on Beijing Air Quality data, with a formal way to keep the "
    "attribution consistent across levels. The third contribution is DyHuCoG, "
    "a Dynamic Hypergraph Cooperative Game for preference-aware "
    "recommendation that brings attribution inside the model itself, "
    "evaluated on MovieLens-1M and Amazon-Book.\n\n"
    "Together, these three contributions follow one progression: from "
    "explanation, to scalability, and finally to action."
)


def main():
    notes = json.load(open(NOTES_IN))
    before = len(notes[str(SLIDE)].split())
    notes[str(SLIDE)] = NOTE_16
    json.dump(notes, open(NOTES_OUT, "w"), ensure_ascii=False, indent=1)

    prs = Presentation(SRC)
    prs.slides[SLIDE - 1].notes_slide.notes_text_frame.text = NOTE_16
    prs.save(DST)

    spoken = sum(len(notes[str(i)].split()) for i in range(1, 77))
    print(f"note {SLIDE}: {before} words -> {len(NOTE_16.split())} words")
    print(f"wrote {os.path.relpath(DST, HERE)}")
    print(f"wrote {os.path.relpath(NOTES_OUT, HERE)}")
    print(f"spoken words slides 1-76: {spoken} "
          f"({spoken / 130:.1f} min @130 wpm, {spoken / 120:.1f} @120 wpm)")
    print(f"\nnote {SLIDE}:\n{NOTE_16}")


if __name__ == "__main__":
    main()
