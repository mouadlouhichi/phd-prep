#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# patch_v31.py: v30 -> v31
#
# One change: the speech on slide 4 is cut to about eighty words, keeping the
# three questions the slide asks and the tension the banner states. The note
# reads as speech, not as prose: short sentences, "First / Second / Third".
# ---------------------------------------------------------------------------
import json
import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v30.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v31.pptx")
NOTES_IN = os.path.join(HERE, "notes_v30.json")
NOTES_OUT = os.path.join(HERE, "notes_v31.json")

SLIDE = 4

NOTE_4 = (
    "Three questions motivate this work.\n\n"
    "First, how do black-box AI systems shape what billions of users see, buy "
    "and watch? Recommenders already shape news, study, health and credit "
    "decisions.\n\n"
    "Second, why do state-of-the-art recommenders and clustering pipelines "
    "stay black boxes?\n\n"
    "Third, how can transparency be built into the model instead of added "
    "afterwards?\n\n"
    "That is the tension: as models gain power, they lose the transparency we "
    "need. This thesis treats accuracy and interpretability as two goals to be "
    "met together."
)


def main():
    notes = json.load(open(NOTES_IN))
    before = len(notes[str(SLIDE)].split())
    notes[str(SLIDE)] = NOTE_4
    after = len(NOTE_4.split())
    json.dump(notes, open(NOTES_OUT, "w"), ensure_ascii=False, indent=1)

    prs = Presentation(SRC)
    prs.slides[SLIDE - 1].notes_slide.notes_text_frame.text = NOTE_4
    prs.save(DST)

    spoken = sum(len(notes[str(i)].split()) for i in range(1, 77))
    print(f"note {SLIDE}: {before} words -> {after} words")
    print(f"wrote {os.path.relpath(DST, HERE)}")
    print(f"wrote {os.path.relpath(NOTES_OUT, HERE)}")
    print(f"spoken words slides 1-76: {spoken} ({spoken / 130:.1f} min @130 wpm, "
          f"{spoken / 120:.1f} @120 wpm)")


if __name__ == "__main__":
    main()
