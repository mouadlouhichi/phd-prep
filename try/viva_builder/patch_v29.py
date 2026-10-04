#!/usr/bin/env python3
# ---------------------------------------------------------------------------
# patch_v29.py: v28 -> v29   -- the reference bands, in full
#
# Every slide that cites something carries a one-line band above the running
# footer, in the short form "[21] Zhang 2020". Two problems with it:
#
#   too short   it is not a reference, it is a pointer
#   hidden      on 12 slides the content was drawn over it (the band sat at
#               z-order 4 of 37, and cards/note bars ending at y = 10.28
#               covered the second half of the band)
#
# This patch rewrites every band with the full citation, taken from the deck's
# own reference list on slides 73-74, laid out at 14 pt over up to three lines,
# drawn last (so nothing can cover it again) and with the content that would
# have covered it moved up, group by group, only as far as needed.
#
# Slides that cite four or more works (23 and 52) would need four or five
# lines; on those the citation keeps authors, venue, volume, pages and year but
# drops the paper title, which is what makes it fit in three lines. Everything
# is measured with the metrics of the font embedded in the file (Nunito
# Semi-Bold), so the line counts are real, not estimates.
# ---------------------------------------------------------------------------
import io
import json
import os
import re
import zipfile

from pptx import Presentation
from pptx.util import Emu

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v28.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v29.pptx")

EMU_IN = 914400
BAND_BOTTOM = 10.54          # the white strip's bottom edge, just above the footer
LINE_PT = 17.0               # the band's existing line spacing
PAD_IN = 0.06
MAX_LINES = 3
SAFETY = 0.97                # measure the wrap on a slightly narrower box than PowerPoint will use
CLEAR = 0.03                 # breathing room between content and the band
FIXED_TOP = 3.40             # everything above this belongs to the header/tab row
FOOTER_TOP = 10.55           # footer text and page number live below this

# Refs whose title is dropped when a slide cites too many works for three full
# lines (see the header comment). Authors, venue, volume, pages and year stay.
COMPACT = {
    "7":  "Ribeiro, M. T., Singh, S. & Guestrin, C. KDD, 1135–1144 (2016).",
    "10": "Koren, Y., Bell, R. & Volinsky, C. Computer 42(8), 30–37 (2009).",
    "11": "He, X. et al. WWW, 173–182 (2017).",
    "12": "He, X. et al. SIGIR, 639–648 (2020).",
    "13": "Xia, L. et al. SIGIR, 70–79 (2022).",
    "14": "Xiangyi, Y. et al. Int. J. Data Science and Analytics 19(2), 269–281 (2025).",
    "15": "Zhang, D. et al. The Web Conference, 3655–3666 (2024).",
    "21": "Zhang, Y. & Chen, X. Foundations and Trends in IR 14(1), 1–101 (2020).",
    "23": "Zhu, Z. et al. KDD, 2439–2449 (2021). · Hurley, N. & Zhang, M. ACM TOIT 10(4) (2011).",
}

SEP = "   "                  # between two references, as in the candidate's example
SUBSEP = "  ·  "             # between two works inside one numbered reference


# ---------------------------------------------------------------------------
# measuring
# ---------------------------------------------------------------------------
def embedded_font(path, typeface):
    from fontTools.ttLib import TTFont
    z = zipfile.ZipFile(path)
    rels = z.read("ppt/_rels/presentation.xml.rels").decode()
    pres = z.read("ppt/presentation.xml").decode()
    rid = None
    for m in re.finditer(r"<p:embeddedFont>(.*?)</p:embeddedFont>", pres, re.S):
        if f'typeface="{typeface}"' in m.group(1):
            rid = re.search(r'r:id="(rId\d+)"', m.group(1)).group(1)
    if rid is None:
        return None
    part = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"', rels).group(1)
    if not part.startswith("ppt/"):
        part = "ppt/" + part.lstrip("/")
    data = z.read(part)
    for off in range(0, 4096):
        if data[off:off + 4] in (b"OTTO", b"\x00\x01\x00\x00"):
            try:
                return TTFont(io.BytesIO(data[off:]), lazy=True)
            except Exception:
                continue
    return None


class Measurer:
    def __init__(self, font):
        self.upm = font["head"].unitsPerEm
        self.cmap = font.getBestCmap()
        self.hmtx = font["hmtx"]

    def width(self, text, size=14.0):
        return sum(self.hmtx[self.cmap.get(ord(c)) or self.cmap.get(ord("?"))][0]
                   / self.upm * size for c in text)

    def lines(self, text, box_pt):
        n, cur = 1, 0.0
        for word in text.split(" "):
            w = self.width(word + (" " if cur else ""))
            if cur + w > box_pt and cur > 0:
                n += 1
                cur = self.width(word)
            else:
                cur += w
        return n


# ---------------------------------------------------------------------------
# geometry helpers
# ---------------------------------------------------------------------------
def rect(sh):
    return (Emu(sh.left).inches, Emu(sh.top).inches,
            Emu(sh.left + sh.width).inches, Emu(sh.top + sh.height).inches)


def inside(inner, outer):
    return (inner[0] >= outer[0] - 0.02 and inner[2] <= outer[2] + 0.02
            and inner[1] >= outer[1] - 0.02 and inner[3] <= outer[3] + 0.02)


def center(sh):
    r = rect(sh)
    return ((r[0] + r[2]) / 2, (r[1] + r[3]) / 2)


def in_rect(pt, r):
    return r[0] - 0.02 <= pt[0] <= r[2] + 0.02 and r[1] - 0.02 <= pt[1] <= r[3] + 0.02


# ---------------------------------------------------------------------------
def main():
    refs = {}
    prs = Presentation(SRC)
    for si in (73, 74):
        shapes = list(prs.slides[si - 1].shapes)
        for j, sh in enumerate(shapes[:-1]):
            if not sh.has_text_frame:
                continue
            m = re.match(r"\[(\d+)\]$", sh.text_frame.text.strip())
            if m:
                refs[m.group(1)] = SUBSEP.join(
                    p.strip() for p in re.split(r"\s+·\s+", shapes[j + 1].text_frame.text.replace("\n", " "))
                    if p.strip())
    report = [f"reference list read from slides 73-74: {len(refs)} entries"]

    meas = Measurer(embedded_font(SRC, "Nunito Semi-Bold"))
    box_pt = 17.76 * 72 * SAFETY      # conservative: a wrap that is borderline at 100 % still counts here

    shifted_total = 0
    for i, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        band = strip = None
        for sh in shapes:
            if not sh.has_text_frame:
                continue
            r = rect(sh)
            if (sh.name.startswith("TextBox") and sh.text_frame.text.strip().startswith("[")
                    and 10.0 <= r[1] <= 10.2 and "] " in sh.text_frame.text[:8]):
                band = sh
        for sh in shapes:
            r = rect(sh)
            if (sh.name == "Rounded Rectangle 4" and not (sh.text_frame.text or "").strip()
                    and 9.9 <= r[1] <= 10.1):
                strip = sh
        if band is None or strip is None:
            continue

        nums = re.findall(r"\[(\d+)\]", band.text_frame.text)
        full = SEP.join(f"[{n}] {refs[n]}" for n in nums)
        lines = meas.lines(full, box_pt)
        form = "full"
        text = full
        if lines > MAX_LINES:
            compact = SEP.join(f"[{n}] {COMPACT.get(n, refs[n])}" for n in nums)
            clines = meas.lines(compact, box_pt)
            if clines < lines:
                text, lines, form = compact, clines, "compact"

        h = lines * LINE_PT / 72 + PAD_IN
        top = BAND_BOTTOM - h

        # --- move whatever would sit under the band, group by group ----------
        content = []
        hdr_bottom = 0.0
        for sh in shapes:
            if sh._element is band._element or sh._element is strip._element:
                continue
            r = rect(sh)
            if r[1] > FOOTER_TOP:
                continue
            if r[3] <= FIXED_TOP:
                hdr_bottom = max(hdr_bottom, r[3])
                continue
            if r[1] < FIXED_TOP:
                hdr_bottom = max(hdr_bottom, r[3])       # tab row that hangs low
            content.append(sh)

        limit = top - CLEAR
        intruders = [sh for sh in content if rect(sh)[3] > limit]
        moved, deltas = [], []

        if intruders:
            # one group = an intruder plus everything drawn inside it
            groups = [[sh] for sh in intruders]
            for sh in content:
                if sh in intruders:
                    continue
                c = center(sh)
                for g in groups:
                    if any(in_rect(c, rect(x)) or inside(rect(sh), rect(x)) for x in g):
                        g.append(sh)
                        break

            for g in groups:
                gt = min(rect(x)[1] for x in g)
                gb = max(rect(x)[3] for x in g)
                needed = gb - limit
                if needed <= 0:
                    continue

                # how far may the group rise before it touches what is above it
                gx0 = min(rect(x)[0] for x in g)
                gx1 = max(rect(x)[2] for x in g)
                ceiling = 0.0
                for sh in content:
                    if sh in g:
                        continue
                    r = rect(sh)
                    if (r[3] <= gt + 0.02 and not (r[2] <= gx0 + 0.02 or r[0] >= gx1 - 0.02)):
                        ceiling = max(ceiling, r[3])
                allowed = max(0.0, gt - max(ceiling + CLEAR, hdr_bottom + 0.02))
                delta = min(needed, allowed)

                for x in g:
                    x.top = Emu(int((Emu(x.top).inches - delta) * EMU_IN))
                moved.extend(g)
                deltas.append(round(delta, 2))

                shortfall = needed - delta
                if shortfall > 0.005:
                    # a card panel can give up the empty strip under its text
                    shaved = []
                    for x in g:
                        r = rect(x)
                        if (str(x.shape_type).startswith("AUTO_SHAPE")
                                and r[3] > limit + 0.001
                                and Emu(x.height).inches - shortfall > 0.6):
                            x.height = Emu(int((Emu(x.height).inches - shortfall) * EMU_IN))
                            shaved.append(x.name)
                    report.append(f"  slide {i}: group lifted {delta:.2f} in (ceiling {ceiling:.2f}), "
                                  f"panels shortened by {shortfall:.2f} in {shaved if shaved else '(none!)'}")
                    if not shaved:
                        report.append(f"  slide {i}: !! group ending {gb:.2f} still "
                                      f"{shortfall:.2f} in below the band")

        # --- the band itself --------------------------------------------------
        tf = band.text_frame
        run = tf.paragraphs[0].runs[0]
        run.text = text
        for p in list(tf.paragraphs)[1:]:
            p._p.getparent().remove(p._p)
        band.left, band.width = Emu(int(1.12 * EMU_IN)), Emu(int(17.76 * EMU_IN))
        band.top, band.height = Emu(int(top * EMU_IN)), Emu(int(h * EMU_IN))
        strip.left, strip.width = band.left, band.width
        strip.top, strip.height = Emu(int((top - 0.04) * EMU_IN)), Emu(int((h + 0.08) * EMU_IN))

        # --- strip and band last in the tree, so nothing can cover them -------
        spTree = slide.shapes._spTree
        for sh in (strip, band):
            spTree.remove(sh._element)
        spTree.append(strip._element)
        spTree.append(band._element)

        shifted_total += len(moved)
        report.append(f"slide {i:>2}: {len(nums)} ref, {form:<7} {lines} lines, "
                      f"band {top:.2f}-{BAND_BOTTOM:.2f}, "
                      f"{len(moved)} shapes moved up by {deltas if deltas else '-'}")

    prs.save(DST)
    report.append(f"\nwrote {os.path.relpath(DST, HERE)}")
    report.append(f"shapes moved in total: {shifted_total}")
    print("\n".join(report))


if __name__ == "__main__":
    main()
