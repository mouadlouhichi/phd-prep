"""v33 -> v34: slide 10 (Collaborative Filtering) gets an accurate figure.

The old picture was a decorative bipartite sketch whose similarity chips
(u1~u2 0.87, u1~u4 0.21) were not derived from anything on the slide.  The
new figure (assets/slide10_collaborative_filtering.png, drawn by
fig_slide10.py) shows the mechanism instead: the user-item rating matrix,
Pearson similarity computed on the co-rated columns i1, i2, i4, the missing
cell filled by the similarity-weighted average of the neighbours' ratings
(5.85 / 1.37 = 4.3), and a strip with the real sparsity of the thesis
datasets (MovieLens-1M 4.47% of cells observed, Amazon-Book 0.06%).

Note 10 gains one spoken paragraph that walks the jury through that figure.
Nothing else on the slide or in the deck moves.
"""
import json
import os

from pptx import Presentation
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v33.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v34.pptx")
FIG = os.path.join(HERE, "assets", "slide10_collaborative_filtering.png")

NOTE_10_ADD = (
    "The figure walks through one prediction: similarity is computed on the columns two users both "
    "rated, and the missing cell is filled by a similarity-weighted average of the neighbours' ratings, "
    "5.85 over 1.37, so 4.3 out of 5. The strip below recalls how sparse the real matrices are: "
    "4.47% and 0.06% observed."
)


def main():
    prs = Presentation(SRC)
    slide = prs.slides[9]

    # --- swap the picture part in place (geometry untouched)
    pic = [sh for sh in slide.shapes if sh.shape_id == 33][0]
    rId = pic._element.blipFill.blip.get(qn("r:embed"))
    part = slide.part.related_part(rId)
    assert part.partname.ext == "png", part.partname
    part._blob = open(FIG, "rb").read()

    # --- extend the spoken note
    tf = slide.notes_slide.notes_text_frame
    tf.add_paragraph().text = NOTE_10_ADD

    prs.save(DST)

    # --- notes json + verification
    back = Presentation(DST)
    nj = {str(i): (s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
          for i, s in enumerate(back.slides, 1)}
    json.dump(nj, open(os.path.join(HERE, "notes_v34.json"), "w"), indent=1, ensure_ascii=False)

    old = Presentation(SRC)
    dnotes = [i for i, (a, b) in enumerate(zip(
        [s.notes_slide.notes_text_frame.text for s in old.slides],
        [s.notes_slide.notes_text_frame.text for s in back.slides]), 1) if a != b]
    dimg = []
    for i, (a, b) in enumerate(zip(old.slides, back.slides), 1):
        ha = [sh.image.sha1 for sh in a.shapes if sh.shape_type == 13]
        hb = [sh.image.sha1 for sh in b.shapes if sh.shape_type == 13]
        if ha != hb:
            dimg.append(i)
    dshp = [i for i, (a, b) in enumerate(zip(old.slides, back.slides), 1)
            if len(a.shapes) != len(b.shapes)]
    words = sum(len(nj[str(i)].split()) for i in range(1, 77))
    print("notes changed:", dnotes)
    print("images changed:", dimg)
    print("shape-count changed:", dshp or "none")
    print("slides:", len(back.slides))
    print("spoken words 1-76:", words, f"= {words/130:.1f} min @130, {words/120:.1f} @120")
    print("note 10 words:", len(nj["10"].split()))


if __name__ == "__main__":
    main()
