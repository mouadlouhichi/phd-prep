"""v38 -> v39: slide 12 figure redrawn in the style of thesis Figure 2.4.

The cropped thesis figure showed its seams, so fig_slide12.py now draws a new,
complete figure in the same visual language: dashed panels, grey node circles,
"Graph:" with the adjacency matrix W on the left, "Hypergraph:" with hyperedge
curves e1 (solid), e2 (dashed), e3 (dash-dot), the caption "Hyperedge group 1"
and the incidence matrix H on the right.  Both matrices are generated from the
edge and hyperedge lists, so the figure is self-consistent.

The picture returns to the full slot (8.75 x 4.88 in, 1:1 at 300 dpi), the
caption names the figure without claiming it is a copy, and note 12 describes
exactly what is on the slide.
"""
import hashlib
import json
import os

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v38.pptx")
DST = os.path.join(HERE, "..", "MOUAD_LOUHICHI_VIVA_40min_BeigeGreen_v39.pptx")
FIG = os.path.join(HERE, "assets", "slide12_graph_vs_hypergraph.png")

CAPTION = "Same data as a graph (W) and as a hypergraph (H)."
NOTE_12 = (
    "The figure redraws the thesis comparison in one frame: on the left, eight nodes as an ordinary "
    "graph with its adjacency matrix W; on the right, the same nodes as a hypergraph, where hyperedges "
    "e1, e2 and e3 each join a subset of nodes and the incidence matrix H records who belongs to which. "
    "Message passing over these structures carries the collaborative signal across hops, and DyHuCoG "
    "makes the hyperedge weights dynamic."
)


def main():
    prs = Presentation(SRC)
    s12 = prs.slides[11]
    pic = [sh for sh in s12.shapes if sh.shape_id == 33][0]
    rId = pic._element.blipFill.blip.get(qn("r:embed"))
    part = s12.part.related_part(rId)
    part._blob = open(FIG, "rb").read()
    pic.left = Inches(10.12)
    pic.top = Inches(4.44)
    pic.width = Inches(8.75)
    pic.height = Inches(4.88)
    cap = [sh for sh in s12.shapes if sh.has_text_frame
           and sh.text_frame.text.startswith("Thesis Figure 2.4")][0]
    cap.text_frame.paragraphs[0].runs[0].text = CAPTION
    s12.notes_slide.notes_text_frame.text = NOTE_12
    prs.save(DST)

    back = Presentation(DST)
    nj = {str(i): (sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else "")
          for i, sl in enumerate(back.slides, 1)}
    json.dump(nj, open(os.path.join(HERE, "notes_v39.json"), "w"), indent=1, ensure_ascii=False)
    old = Presentation(SRC)
    dnotes = [i for i, (a, b) in enumerate(zip(
        [s.notes_slide.notes_text_frame.text for s in old.slides],
        [s.notes_slide.notes_text_frame.text for s in back.slides]), 1) if a != b]
    dimg = [i for i, (a, b) in enumerate(zip(old.slides, back.slides), 1)
            if [sh.image.sha1 for sh in a.shapes if sh.shape_type == 13] !=
               [sh.image.sha1 for sh in b.shapes if sh.shape_type == 13]]
    dshp = [i for i, (a, b) in enumerate(zip(old.slides, back.slides), 1)
            if len(a.shapes) != len(b.shapes)]
    words = sum(len(nj[str(i)].split()) for i in range(1, 76))
    a = hashlib.sha1(open(FIG, "rb").read()).hexdigest()
    pic2 = [sh for sh in back.slides[11].shapes if sh.shape_id == 33][0]
    print("notes changed:", dnotes, "| images:", dimg, "| shapes:", dshp or "none")
    print("fig embedded 1:1:", pic2.image.sha1 == a, pic2.image.size)
    print("slides:", len(back.slides))
    print("spoken words 1-75:", words, f"= {words/130:.1f} min @130, {words/120:.1f} @120")
    print("note 12 words:", len(nj["12"].split()))


if __name__ == "__main__":
    main()
