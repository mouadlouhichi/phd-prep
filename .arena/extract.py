#!/usr/bin/env python3
"""Dump slide titles, shape texts and speaker notes from a pptx to JSON."""
import json, sys
from pptx import Presentation

def extract(path):
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame:
                t = sh.text_frame.text.strip()
                if t:
                    texts.append(t)
            if sh.has_table:
                for row in sh.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append({"n": i, "texts": texts, "notes": notes})
    return out

if __name__ == "__main__":
    data = extract(sys.argv[1])
    json.dump(data, open(sys.argv[2], "w"), ensure_ascii=False, indent=1)
    print("wrote", sys.argv[2], len(data), "slides")
